# -*- coding: utf-8 -*-
"""
按《实验报告模板.docx》的封面与章节结构，生成《时令志 · 二十四节气》Vue 网站的完整实验报告。

用法：
    python build_report.py

要点：
    1. 直接以模板为底稿（Document(模板路径)），封面文字、字体、页码等原样保留；
    2. 在七个 Heading 1 标题后面"插入"正文段落 / 表格 / 图片（用 addnext 移动 XML 节点），
       而不是简单追加到文末；
    3. 正文统一：中文宋体、西文 Times New Roman、小四(12pt)、1.5 倍行距、首行缩进 2 字符；
       代码统一 Consolas 9pt + 浅灰底纹。

注意：本文件里所有"给人看的中文长文本"都用三引号 ''' 包裹，
      这样文本内部可以直接使用英文双引号而不会和 Python 语法冲突。
"""
import os
import shutil
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = r"C:\Users\沙季\Desktop\实验报告模板.docx"
OUT_DOCX = os.path.join(HERE, "叶子阳-20244081130-Vue网站实验报告.docx")
DESKTOP_COPY = r"C:\Users\沙季\Desktop\叶子阳-20244081130-时令志Vue网站实验报告.docx"
SHOTS = os.path.join(HERE, "screenshots")

BODY_ASCII = "Times New Roman"
BODY_EA = "宋体"
CODE_FONT = "Consolas"
CODE_FILL = "F5F6F7"


# --------------------------------------------------------------------------
# 基础工具
# --------------------------------------------------------------------------
def style_run(run, size=12, bold=False, italic=False, ascii_font=BODY_ASCII,
              ea_font=BODY_EA, color=None):
    '''统一设置一段文字的字体（西文与中文 eastAsia 都要设，否则中文不生效）'''
    run.font.size = Pt(size)
    run.font.name = ascii_font
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), ea_font)
    return run


def shade(paragraph, fill=CODE_FILL):
    '''给段落加底纹（用于代码块）'''
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def new_paragraph_after(anchor):
    '''在 anchor 段落之后插入一个新的空段落并返回它（不追加到文末）'''
    p = OxmlElement("w:p")
    anchor._p.addnext(p)
    return Paragraph(p, anchor._parent)


def add_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


class Inserter:
    '''游标：始终记住"最后插入的那一段"，下一个元素就插在它后面'''

    def __init__(self, doc, cursor):
        self.doc = doc
        self.cursor = cursor

    def para(self, text="", size=12, bold=False, indent=True, align=None,
             space_after=6, color=None, line=1.5):
        p = new_paragraph_after(self.cursor)
        pf = p.paragraph_format
        pf.space_after = Pt(space_after)
        pf.space_before = Pt(0)
        pf.line_spacing = line
        if indent:
            pf.first_line_indent = Pt(size * 2)   # 首行缩进 2 个字符
        if align is not None:
            p.alignment = align
        if text:
            style_run(p.add_run(text), size=size, bold=bold, color=color)
        self.cursor = p
        return p

    def heading2(self, text):
        p = self.para(text, size=12, bold=True, indent=False, space_after=4)
        p.paragraph_format.space_before = Pt(10)
        return p

    def numbered(self, index, text, size=12, marker=None):
        '''带编号的条目；marker 给定时用自定义符号（如项目符号）'''
        p = new_paragraph_after(self.cursor)
        pf = p.paragraph_format
        pf.space_after = Pt(3)
        pf.line_spacing = 1.5
        pf.left_indent = Cm(0.95)
        pf.first_line_indent = Cm(-0.55)
        label = marker if marker is not None else str(index)
        style_run(p.add_run(label + ". "), size=size)
        if text:
            style_run(p.add_run(text), size=size)
        self.cursor = p
        return p

    def code(self, code_text, caption=None):
        '''代码块：Consolas 9pt + 浅灰底纹，段内用软换行，保证不被拆成多段'''
        if caption:
            self.para(caption, size=10.5, indent=False, space_after=2, color="595959")
        p = new_paragraph_after(self.cursor)
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(10)
        pf.line_spacing = 1.15
        pf.left_indent = Cm(0.4)
        shade(p)
        lines = code_text.strip("\n").split("\n")
        for i, line in enumerate(lines):
            run = p.add_run(line)
            style_run(run, size=9, ascii_font=CODE_FONT, ea_font=CODE_FONT)
            if i < len(lines) - 1:
                run.add_break()
        self.cursor = p
        return p

    def table(self, rows, widths=None, font_size=10.5):
        '''插入表格：先用 doc.add_table 建好，再把 XML 节点搬到游标后面'''
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        try:
            t.style = self.doc.styles["Table Grid"]
        except KeyError:
            add_table_borders(t)
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.25
                style_run(p.add_run(str(value)), size=font_size, bold=(i == 0))
        if widths:
            for row in t.rows:
                for idx, w in enumerate(widths):
                    row.cells[idx].width = Cm(w)
        self.cursor._p.addnext(t._tbl)
        spacer = new_paragraph_after(Paragraph(t._tbl, self.cursor._parent))
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.line_spacing = 1.0
        self.cursor = spacer
        return t

    def picture(self, filename, caption, width_cm=15.2):
        path = os.path.join(SHOTS, filename)
        p = new_paragraph_after(self.cursor)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Cm(width_cm))
        self.cursor = p
        if caption:
            c = new_paragraph_after(self.cursor)
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.space_after = Pt(10)
            c.paragraph_format.line_spacing = 1.2
            style_run(c.add_run(caption), size=9, color="595959")
            self.cursor = c


def find_heading(doc, text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise SystemExit("模板里找不到标题：" + text)


# --------------------------------------------------------------------------
# 正文内容
# --------------------------------------------------------------------------
def fill(doc):
    # ======================= 一、实验目的 =======================
    h = find_heading(doc, "一、实验目的")
    h.paragraph_format.page_break_before = True
    ins = Inserter(doc, h)

    ins.para(
        '''本次实验的题目是"任选主题，使用 Vue 实现一个网站（只需前端）"。我选择"二十四节气"作为主题，'''
        '''完整实现了一个名为《时令志 · 二十四节气》的纯前端单页应用，并把它部署到了公网。'''
        '''通过这次实验，希望达到以下目的：'''
    )
    purposes = [
        '''掌握 Vue 3 单文件组件（SFC）的编写方式，理解 <template>、<script setup>、<style scoped>'''
        ''' 三部分各自的作用，以及"为什么要用单文件组件，而不是在 HTML 里写一大段脚本"。''',
        '''掌握组合式 API 的核心：响应式基础（ref / reactive）、计算属性 computed、侦听器 watch、'''
        '''生命周期钩子，并学会把可复用的响应式逻辑抽成"组合式函数"（Composables）。''',
        '''掌握组件之间的通信：父传子 props、子传父 emits、组件上的 v-model（defineModel）、透传 '''
        '''Attributes，以及插槽的四种用法（默认插槽、具名插槽、作用域插槽、动态插槽名）、'''
        '''依赖注入 provide / inject。''',
        '''掌握 Vue Router：动态路由参数、路由懒加载、别名与重定向、404 通配符路由、全局导航守卫、'''
        '''滚动行为，并用它把网站组织成 6 个页面的多页应用。''',
        '''掌握 Pinia 状态管理：用 Setup Store 的写法管理主题、收藏与全局提示条等跨页面共享状态，'''
        '''并持久化到浏览器 localStorage。''',
        '''掌握 Vue 的内置进阶能力：Transition / TransitionGroup 过渡动画、Teleport 传送门、'''
        '''KeepAlive 组件缓存、异步组件 defineAsyncComponent 与 Suspense、自定义指令、渲染函数 h()、'''
        '''插件机制与全局错误处理。''',
        '''走完"工程搭建 → 编码 → 构建 → 验证 → 公网部署"的完整闭环，并学会用构建产物和无头浏览器'''
        '''客观验证结果，而不是"看着差不多就行"。''',
    ]
    for i, item in enumerate(purposes, 1):
        ins.numbered(i, item)

    # ======================= 二、实验环境 =======================
    h = find_heading(doc, "二、实验环境")
    ins = Inserter(doc, h)
    ins.para("本次实验使用的软硬件与依赖版本如下表所示。")
    ins.table([
        ["项目", "版本 / 说明"],
        ["操作系统", "Windows 11"],
        ["Node.js", "v24.21.0"],
        ["包管理器", "pnpm 11.28.5"],
        ["核心框架", "Vue 3.5.44（组合式 API + script setup 单文件组件）"],
        ["构建工具", "Vite 8.3.5、@vitejs/plugin-vue 6.0.10"],
        ["前端路由", "Vue Router 5.4.0（HTML5 History 模式）"],
        ["状态管理", "Pinia 4.1.0（Setup Store 写法）"],
        ["样式方案", "原生 CSS + CSS 变量（不引入任何 UI 框架）"],
        ["开发与调试", "VS Code、Microsoft Edge 无头模式（dump-dom / screenshot）"],
        ["版本管理", "Git 2.55.0 + GitHub"],
        ["部署", "GitHub Pages（GitHub Actions 自动构建），同时保留 Netlify 配置"],
        ["数据来源", "自建二十四节气数据集：24 条记录，含三候、习俗、食材、养生、诗词"],
    ], widths=[4.2, 11.0])

    # ======================= 三、实验步骤 =======================
    h = find_heading(doc, "三、实验步骤")
    ins = Inserter(doc, h)

    steps = [
        ("选题与需求分析",
         '''选定"二十四节气"作为主题：数据量适中、文化属性强，天然适合做"列表页 + 详情页 + 交互页"'''
         '''三类页面。最终规划 6 个页面（首页、节气图鉴、节气详情、时令工坊、关于本站、404）和 '''
         '''40 个单文件组件，并确定不使用 UI 框架、不依赖后端，用 CSS 变量、渐变和 emoji 保证'''
         '''视觉统一且完全离线可用。'''),
        ("搭建工程骨架",
         '''执行 pnpm add vue vue-router pinia 与 pnpm add -D vite @vitejs/plugin-vue 安装依赖；'''
         '''编写 package.json 的 dev / build / preview 脚本；在 vite.config.js 中配置 @ 别名指向 src，'''
         '''并把 base 做成可由环境变量 BASE_PATH 控制（为了后面部署到子路径）；index.html 只保留'''
         '''一个 id 为 app 的 div 作为挂载点。'''),
        ("编写入口与全局装配",
         '''在 src/main.js 中按顺序完成：createApp(App) 创建应用 → app.use(createPinia()) 装状态管理 → '''
         '''app.use(router) 装路由 → app.use(SiteInfoPlugin) 装自定义插件 → app.directive 注册全局指令 → '''
         '''app.config.errorHandler 设置全局错误兜底 → app.mount('#app') 挂载。'''),
        ("设计数据层",
         '''把 24 个节气整理成 src/data/terms.js 中的数组，每条记录包含 id、名称、拼音、次序、季节、'''
         '''常年日期、主题渐变色、一句话物候、简介、三候、习俗、时令食材、养生提示、对应诗词。'''
         '''同时导出纯函数 termDate（算某年的公历日期）、getCurrentTerm、getNextTerm、formatDate，'''
         '''供组件在 computed 里直接调用——把"数据"和"视图"彻底分开。'''),
        ("抽取组合式函数",
         '''把可复用的响应式逻辑抽到 src/composables/ 下，共 5 个：useLocalStorage（ref 与 localStorage '''
         '''双向同步）、useCountdown（倒计时，内部管理定时器的创建与销毁）、useSeasonNow（当前与下一个'''
         '''节气，跨天自动刷新）、useTermFilter（搜索 + 季节筛选 + 排序 + 数量统计）、useToast（包装 '''
         '''Pinia 的 ui store，提供 success / error / info 快捷方法）。'''),
        ("实现组件",
         '''公共组件 11 个（BaseCard、BaseButton、BaseModal、BaseTabs、SearchInput、LoadingSpinner、'''
         '''EmptyState、StarRating、ProgressRing、ToastHost、BackToTop），业务组件 22 个（节气 7 个、'''
         '''首页 4 个、工坊 8 个、布局 3 个），页面视图 6 个，加上根组件 App.vue，共 40 个单文件组件。'''
         '''每个组件只做一件事，通过 props、emits 和插槽组合起来。'''),
        ("配置路由与页面",
         '''在 src/router/index.js 中定义 6 条路由和 1 条重定向，页面组件全部用 () => import() 懒加载；'''
         '''用路由的 meta.title 配合全局前置守卫 beforeEach 统一修改浏览器标签标题；用 scrollBehavior '''
         '''处理锚点跳转和"后退时恢复滚动位置"。'''),
        ("接入 Pinia 状态管理",
         '''三个 store：theme（浅色 / 深色 / 跟随系统，结果写到 html 的 data-theme 属性，由 CSS 变量'''
         '''接管配色）、favorites（收藏的节气，持久化到 localStorage）、ui（全局提示条 Toast 与移动端'''
         '''菜单）。跨页面共享"收藏"是选择 Pinia 而不是 props 的根本原因。'''),
        ("增强交互与动画",
         '''页面切换用 Transition；列表增删用 TransitionGroup（含 list-move 重排动画）；弹窗和提示条用 '''
         '''Teleport 渲染到 body；节气圆盘用三角函数把 24 个节气均匀分布到圆周上；自定义指令 v-reveal '''
         '''实现滚动进入视口才浮现，v-highlight 实现关键词高亮。'''),
        ("构建、验证与部署",
         '''执行 pnpm build 产出 dist 目录；用 vite preview 起本地服务，再用 Edge 的无头模式对 6 条路由'''
         '''逐一抓取渲染后的 DOM 并截图，检查关键内容是否存在、有没有 Vue 警告或运行时错误。最后把代码'''
         '''推送到 GitHub 仓库 shiling-zhi，把 Pages 的构建来源设为 GitHub Actions，由 '''
         '''.github/workflows/deploy.yml 自动完成"安装依赖 → 构建 → 生成 404 回退页 → 发布"。'''),
    ]
    for i, (title, body) in enumerate(steps, 1):
        p = ins.numbered(i, "")
        p.runs[0].text = str(i) + ". " + title + "　"
        style_run(p.add_run(body), size=12)

    # ======================= 四、代码示例 =======================
    h = find_heading(doc, "四、代码示例")
    ins = Inserter(doc, h)
    ins.para("下面按知识点摘录几段核心代码，并说明它解决的问题。")

    ins.heading2("1. 组件上的 v-model：让评分组件支持双向绑定")
    ins.para(
        '''父组件只要写 <StarRating v-model="score" />。子组件用 defineModel() 声明一个可写的 model，'''
        '''读它等于读 prop，写它会自动向父组件 emit update:modelValue，省掉了以前"传 prop + 监听事件 + '''
        '''手动 emit"三件套。'''
    )
    ins.code('''
// components/common/StarRating.vue
const model = defineModel({ type: Number, default: 0 })   // 等价于声明 prop + 事件
const hoverIndex = ref(0)
const displayValue = computed(() => hoverIndex.value || model.value)

function setScore(value) {
  if (props.readonly) return
  model.value = value === model.value ? 0 : value      // 再点一次同样的分数就清零
  emit('change', model.value)
}
''')
    ins.para(
        '''在详情页里，这个 model 甚至不是普通的 ref，而是 computed 的 getter/setter——读取时从 '''
        '''localStorage 里取该节气的分数，写入时把新分数合并回对象，做到了"组件不用关心数据存在哪里"。'''
    )

    ins.heading2("2. 一个组件多个 v-model：筛选工具条")
    ins.para(
        '''图鉴页的筛选条件有 4 个（关键词、季节、排序、是否只看收藏）。用 v-model:xxx 的写法可以一次'''
        '''绑定多个，组件内部对应 4 个 defineModel，父组件的模板非常直观。'''
    )
    ins.code('''
<!-- 父组件：views/TermsView.vue -->
<TermFilterBar
  v-model:keyword="keyword"
  v-model:season="season"
  v-model:sortBy="sortBy"
  v-model:onlyFavorites="onlyFavorites"
  :counts="counts" :total="visibleTerms.length" @reset="reset" />

// 子组件：components/terms/TermFilterBar.vue
const keyword       = defineModel('keyword',       { type: String,  default: '' })
const season        = defineModel('season',        { type: String,  default: 'all' })
const sortBy        = defineModel('sortBy',        { type: String,  default: 'order' })
const onlyFavorites = defineModel('onlyFavorites', { type: Boolean, default: false })
''')

    ins.heading2("3. 组合式函数：把倒计时逻辑封装成一行调用")
    ins.para(
        '''倒计时需要每秒刷新、需要在卸载时清理定时器、还需要在目标变化时重新计算。把这些都放进 '''
        '''useCountdown，组件里只要一行解构就能用，而且生命周期由组合式函数自己管理。'''
    )
    ins.code('''
// composables/useCountdown.js（节选）
export function useCountdown(target) {
  const now = ref(Date.now())
  let timer = null
  const targetDate  = computed(() => resolveTarget(target))
  const remainingMs = computed(() => Math.max(0, targetDate.value.getTime() - now.value))
  const days        = computed(() => Math.floor(remainingMs.value / 86400000))
  const isOver      = computed(() => remainingMs.value <= 0)

  function start() { stop(); timer = window.setInterval(() => (now.value = Date.now()), 1000) }
  function stop()  { if (timer) { window.clearInterval(timer); timer = null } }

  watch(targetDate, () => (now.value = Date.now()))
  onMounted(start)        // 组合式函数里也能用生命周期钩子
  onUnmounted(stop)       // 一定要清理，否则组件销毁后定时器还在跑
  return { days, isOver, start, stop }
}
''')

    ins.heading2("4. Pinia：跨页面共享的收藏状态")
    ins.para(
        '''首页、图鉴页、详情页都要读同一份收藏数据，所以放进 store 而不是层层传 props。这里用的是 '''
        '''Setup Store 写法：state 用 ref、getters 用 computed、actions 用普通函数。'''
    )
    ins.code('''
// stores/favorites.js（节选）
export const useFavoritesStore = defineStore('favorites', () => {
  const ids   = useLocalStorage('shiling:favorites', [])   // state，自动持久化
  const count = computed(() => ids.value.length)           // getter
  const list  = computed(() => ids.value.map(id => TERM_MAP[id]).filter(Boolean)
                                   .sort((a, b) => a.order - b.order))
  function toggle(id) {                                    // action
    const i = ids.value.indexOf(id)
    if (i > -1) { ids.value.splice(i, 1); return false }
    ids.value.push(id); return true
  }
  return { ids, count, list, toggle }
})

// 组件中使用：必须用 storeToRefs，直接解构会丢失响应性
const favorites = useFavoritesStore()
const { ids: favoriteIds, count } = storeToRefs(favorites)
''')

    ins.heading2("5. 动态具名插槽：一个 Tabs 组件容纳任意面板")
    ins.para(
        '''详情页有"三候 / 习俗 / 时令 / 养生"四个面板。BaseTabs 把当前选中的 tab 名当作插槽名'''
        '''（slot 的 name 用 v-bind 动态绑定），父组件按名字把内容填进去，组件本身完全不知道有几块内容。'''
    )
    ins.code('''
<!-- components/common/BaseTabs.vue -->
<Transition name="fade" mode="out-in">
  <div :key="active">
    <slot :name="active">{{ activeTab?.emptyText || '暂无内容' }}</slot>
  </div>
</Transition>

<!-- 使用方 components/terms/TermDetailPanel.vue -->
<BaseTabs v-model="activeTab" :tabs="tabs">
  <template #hou>     <ol class="hou"> ... </ol>           </template>
  <template #customs> <ul class="chips"> ... </ul>         </template>
  <template #foods>   <ul class="chips is-food"> ... </ul> </template>
  <template #health>  <p class="health"> ... </p>          </template>
</BaseTabs>
''')

    ins.heading2("6. 异步组件 + Suspense：按需下载 + 加载占位")
    ins.para(
        '''冷知识卡片的内容组件里写了顶层 await（模拟请求后端接口）。它由 defineAsyncComponent 包装，'''
        '''外层再用 Suspense 兜住：等待期间显示 fallback，加载失败显示 errorComponent；打包后它还是'''
        '''一个独立的 JS 文件，用到时才下载。'''
    )
    ins.code('''
// components/workshop/AsyncFactCard.vue（节选）
const AsyncFactContent = defineAsyncComponent({
  loader: () => import('./AsyncFactContent.vue'),   // 独立 chunk，按需加载
  loadingComponent: LoadingBox,
  errorComponent: ErrorBox,
  delay: 200,        // 200ms 内加载完就不显示 loading，避免闪烁
  timeout: 8000
})

<!-- 模板：Suspense 的 fallback 负责"等待中的样子" -->
<Suspense>
  <AsyncFactContent :key="reloadKey" />
  <template #fallback>
    <LoadingSpinner label="加载中…（Suspense fallback）" />
  </template>
</Suspense>
''')

    ins.heading2("7. 自定义指令：滚动进入视口才浮现")
    ins.para(
        '''指令适合封装"直接操作 DOM"的逻辑。v-reveal 用 IntersectionObserver 监听元素是否进入视口，'''
        '''进入后加上 is-visible 类，配合 CSS 过渡做出滚动浮现的效果；元素卸载时断开观察，避免内存泄漏。'''
    )
    ins.code('''
// directives/reveal.js（节选）
export const vReveal = {
  mounted(el, binding) {
    const { delay = 0, once = true, threshold = 0.16 } = binding.value || {}
    el.classList.add('reveal')
    const observer = new IntersectionObserver((entries) => {
      entries.forEach((e) => {
        if (e.isIntersecting) {
          window.setTimeout(() => el.classList.add('is-visible'), delay)
          if (once) observer.unobserve(el)
        } else if (!once) el.classList.remove('is-visible')
      })
    }, { threshold })
    observer.observe(el)
    el._revealObserver = observer        // 存起来，卸载时才能断开
  },
  unmounted(el) { el._revealObserver?.disconnect(); delete el._revealObserver }
}
''')

    ins.heading2("8. 路由：动态参数、懒加载与守卫")
    ins.code('''
// router/index.js（节选）
const routes = [
  { path: '/',          name: 'home',        component: () => import('@/views/HomeView.vue') },
  { path: '/terms',     name: 'terms',       component: () => import('@/views/TermsView.vue') },
  // :id 是动态参数；props: true 把它直接变成组件的 prop，组件不必再 useRoute()
  { path: '/terms/:id', name: 'term-detail', component: () => import('@/views/TermDetailView.vue'),
    props: true },
  { path: '/about',     name: 'about',       component: () => import('@/views/AboutView.vue'),
    alias: ['/about-us'] },
  { path: '/list',      redirect: { name: 'terms' } },
  { path: '/:pathMatch(.*)*', name: 'not-found',
    component: () => import('@/views/NotFoundView.vue') }
]

// 全局前置守卫：统一维护页面标题
router.beforeEach((to) => {
  document.title = to.meta?.title
    ? to.meta.title + ' | 时令志 · 二十四节气'
    : '时令志 · 二十四节气'
  return true
})
''')

    # ======================= 五、实验结果 =======================
    h = find_heading(doc, "五、实验结果")
    ins = Inserter(doc, h)
    ins.para('''网站已完整实现并成功部署到公网，本地与线上均验证通过。''')
    ins.table([
        ["项目", "结果"],
        ["公网访问地址", "https://aysakili.github.io/shiling-zhi/"],
        ["源码仓库", "https://github.com/AySakili/shiling-zhi"],
        ["页面数量", "6 个（首页 / 节气图鉴 / 节气详情 / 时令工坊 / 关于本站 / 404）"],
        ["单文件组件", "40 个（公共 11 + 节气 7 + 首页 4 + 工坊 8 + 布局 3 + 页面 6 + 根组件 1）"],
        ["可复用逻辑", "组合式函数 5 个、Pinia store 3 个、自定义指令 2 个、插件 1 个"],
        ["代码规模", "src 下 .vue 文件 6223 行、.js 文件 1231 行，合计约 7454 行"],
        ["构建产物", "dist 共 32 个文件、约 277 KB（JS 221.4 KB、CSS 54.3 KB，已按页面分包）"],
        ["数据规模", "24 条节气记录，含三候 72 条、习俗、食材、养生、诗词等字段"],
        ["部署方式", "GitHub Actions 自动构建并发布到 GitHub Pages"],
    ], widths=[3.6, 11.6])

    ins.heading2("5.1 各页面运行效果")
    ins.picture("01-home.png", "图 1　首页：当前节气（寒露）、下一个节气倒计时、二十四节气圆盘")
    ins.picture("02-terms.png", "图 2　节气图鉴：搜索、季节筛选、排序与卡片列表")
    ins.picture("03-detail.png", "图 3　节气详情：动态路由 /terms/shuangjiang 与三候面板")
    ins.picture("04-workshop.png", "图 4　时令工坊：动态组件切换与时令食材推荐器")
    ins.picture("05-about.png", "图 5　关于本站：技术栈与知识点—文件对照表")
    ins.picture("06-notfound.png", "图 6　404 页面：通配符路由与随机推荐")
    ins.picture("07-wheel.png", "图 7　二十四节气圆盘特写（24 个节气按三角函数均匀分布）")

    ins.heading2("5.2 功能验证")
    checks = [
        '''构建验证：pnpm build 成功，117 个模块被转换，输出按页面拆分的多个 chunk，'''
        '''说明路由懒加载生效。''',
        '''渲染验证：用 Microsoft Edge 无头模式对 6 条路由逐一抓取渲染后的 DOM，确认首页出现'''
        '''"今日 寒露"、圆盘渲染出 24 个节气按钮、详情页三候内容正常显示、404 文案正确，'''
        '''且全文没有出现 Vue 警告或 TypeError。''',
        '''部署验证：GitHub Actions 工作流执行"安装依赖 → 构建（base 设为 /shiling-zhi/）→ '''
        '''生成 404.html → 发布"全部成功，公网地址可正常访问。''',
        '''交互验证：收藏在首页与图鉴页之间同步；主题切换后刷新仍然保持；筛选条件在跳转详情页再'''
        '''返回后依然保留（KeepAlive 生效）；倒计时每秒更新。''',
    ]
    for i, c in enumerate(checks, 1):
        ins.numbered(i, c)

    # ======================= 六、问题与解决 =======================
    h = find_heading(doc, "六、问题与解决")
    ins = Inserter(doc, h)
    ins.para("开发过程中遇到了不少问题，下面记录其中印象最深的几个。")

    problems = [
        ("节气圆盘的 24 个点全部叠到了圆心",
         '''最初给圆环加了一个缓慢自转的 CSS 动画，同时用 transform: translate(x, y) 给每个圆点定位。'''
         '''结果所有圆点都挤在圆心，只剩下一团白色。原因是：CSS 动画里的 transform 会覆盖行内样式设置的 '''
         '''transform，两者写在同一个属性上互相冲突。''',
         '''把定位方式从 transform 改成 left / top（用 calc(50% 加上偏移量) 计算），动画只保留在装饰层上；'''
         '''同时用 ResizeObserver 监听容器宽度，动态计算可用半径，让圆盘在手机上也能自适应。'''),
        ("运行时组件里写 template 字符串报错",
         '''异步组件的 loadingComponent / errorComponent 一开始图省事写成了带 template 字符串的对象。'''
         '''运行时控制台报错：Template compilation is not supported in this build。''',
         '''原因是 Vite 默认使用 Vue 的"运行时版"，不含模板编译器，template 必须在构建时被编译。'''
         '''改成用渲染函数 h() 写这两个小组件即可：setup() { return () => h(LoadingSpinner, {...}) }。'''),
        ("Vite 8 不再接受对象形式的 manualChunks",
         '''在 vite.config.js 里按老习惯写了 output.manualChunks = { vue: [...] }，构建时提示'''
         '''Invalid type: Expected Function but received Object。''',
         '''Vite 8 底层换成了 rolldown，分包配置方式有变化。而本项目的页面本来就是懒加载，默认分包'''
         '''已经很合理，于是直接去掉手写分包，交给构建工具自动处理。'''),
        ("单页应用刷新子路由出现 404",
         '''直接访问或刷新 /terms/shuangjiang 时，服务器会去找这个物理路径，自然找不到。''',
         '''这是 history 模式路由的固有代价，需要服务器配合：Netlify 侧写 public/_redirects 与 '''
         '''netlify.toml 把任意路径回退到 index.html；GitHub Pages 侧没有回退配置，于是在工作流里把 '''
         '''dist/index.html 复制一份成 404.html，让找不到的路径返回它，再由前端路由接管。'''),
        ("部署到子路径后 JS 与 CSS 全部 404",
         '''GitHub Pages 的项目站点地址形如 https://用户名.github.io/仓库名/，而默认构建产物引用的是 '''
         '''/assets/...，少了一层仓库名。''',
         '''把 vite.config.js 的 base 改成由环境变量 BASE_PATH 控制，在 GitHub Actions 里注入'''
         '''仓库名对应的子路径，本地和 Netlify 则保持根路径。路由的 '''
         '''createWebHistory(import.meta.env.BASE_URL) 会自动跟着变。'''),
        ("KeepAlive 缓存的页面里 onMounted 不再触发",
         '''图鉴页被 KeepAlive 缓存后，从详情页返回时组件没有重新创建，onMounted 里"读取 URL 查询参数"'''
         '''的逻辑就不再执行，从页脚点"秋季"进不去筛选状态。''',
         '''改用 onActivated 钩子（缓存组件每次被激活都会触发），并额外监听 route.query 的变化，'''
         '''处理同一页面内查询参数变化的情况。'''),
        ("Pinia 解构后数据不再响应",
         '''写成 const { list } = useFavoritesStore() 之后，收藏变化了页面却不更新。''',
         '''直接解构 store 会丢掉响应性，必须用 storeToRefs() 解构 state 与 getter（方法仍从 store 上调用）。'''),
        ("定时器与全局事件监听的清理",
         '''倒计时、滚动监听、键盘方向键切换节气等地方都注册了副作用，若不清理会造成内存泄漏，'''
         '''甚至在组件销毁后仍然修改已不存在的状态。''',
         '''统一在 onUnmounted 中清理：clearInterval、removeEventListener、observer.disconnect()；'''
         '''组合式函数 useCountdown 内部自己完成清理，做到"谁申请、谁释放"。'''),
        ("同一元素上写了两个同名自定义指令",
         '''演示 v-highlight 的参数与修饰符时，本想在一个元素上同时写两次 v-highlight，结果编译报错。''',
         '''Vue 中一个元素上同名指令只能出现一次。改成两个元素分别演示不同写法，顺便把差异展示得更清楚。'''),
    ]
    for i, (title, cause, solution) in enumerate(problems, 1):
        p = ins.numbered(i, "")
        p.runs[0].text = str(i) + ". " + title + "　"
        style_run(p.add_run(cause), size=12)
        q = ins.para("", indent=False)
        q.paragraph_format.left_indent = Cm(0.95)
        style_run(q.add_run("解决办法："), size=12, bold=True)
        style_run(q.add_run(solution), size=12)

    # ======================= 七、实验总结 =======================
    h = find_heading(doc, "七、实验总结")
    ins = Inserter(doc, h)
    ins.para(
        '''这次实验从选题、搭工程、写组件、接路由和状态管理，一直到打包部署上线，完整地走了一遍'''
        '''现代前端项目的流程。最后交出的成果是：6 个页面、40 个单文件组件、约 7400 行代码，'''
        '''构建产物 277 KB，已经可以通过公网地址访问。'''
    )
    ins.heading2("主要收获")
    gains = [
        '''真正理解了"响应式"是怎么工作的：用 ref 声明状态，computed 会缓存并自动追踪依赖变化，'''
        '''watch 用来处理副作用（写 localStorage、打日志、重启定时器）。一开始我习惯把什么都写成'''
        '''状态再手动同步，写完之后才体会到"派生的数据不要手动维护"有多省事——比如图鉴页的筛选结果、'''
        '''数量统计、是否处于筛选状态，全部由 computed 推导，代码里找不到一处"手动更新列表"的逻辑。''',
        '''组件设计的核心是"职责单一 + 通过接口通信"。全站的组件分成三层：基础组件（按钮、卡片、弹窗）、'''
        '''业务组件（节气卡片、圆盘、推荐器）、页面视图。基础组件只认 props 和插槽，完全不知道"节气"'''
        '''是什么，因此可以在任何页面复用；页面负责把数据取好再往下传，组件树因此很浅，'''
        '''出了问题很容易定位。''',
        '''学会了用"验证"代替"感觉"：以前写完页面刷新看看没问题就算完成。这次用无头浏览器抓取渲染后的 '''
        '''DOM 并扫描错误关键字，能发现肉眼看不到的问题；用构建产物清单可以确认代码分割是否真的生效。'''
        '''这套做法在答辩时也很有说服力——每个结论都有对应的输出证据。''',
        '''工程化不再是"会敲启动命令"：环境变量控制部署路径、CI 自动构建发布、单页应用的路由回退、'''
        '''缓存策略、依赖版本差异，这些"和业务无关"的东西，恰恰决定了项目能不能真正上线。''',
    ]
    for i, t in enumerate(gains, 1):
        ins.numbered(i, t)

    ins.heading2("不足与改进方向")
    todos = [
        '''数据是写死的静态数据，日期用的是常年日期。下一步可以接入公开的节气接口，或在构建期用天文'''
        '''算法算出精确时刻，并加上"当前处于第几候"的推算。''',
        '''目前没有任何自动化测试。可以对 useTermFilter、getNextTerm 这类纯函数与组合式函数补充单元测试，'''
        '''对关键交互补充组件测试，避免以后改代码时破坏已有功能。''',
        '''首屏仍是客户端渲染，对搜索引擎不友好。可以考虑用 Nuxt 或预渲染方案做服务端渲染，'''
        '''同时改善首屏白屏时间。''',
        '''无障碍与国际化还有提升空间：虽然已经加了 aria 标签和键盘操作（方向键切换节气、Esc 关闭弹窗），'''
        '''但对比度、焦点顺序还没有系统检查；界面文案也全部硬编码为中文，可以引入 i18n。''',
    ]
    for i, t in enumerate(todos, 1):
        ins.numbered(i, t)

    ins.para(
        '''总的来说，这次实验让我把 Vue 从"能看懂"推进到了"能自己设计"。最大的体会是：框架提供的 API '''
        '''只是工具，真正决定代码质量的，是"哪些状态该被派生、哪些逻辑该被封装、组件之间的边界画在哪里"'''
        '''这些设计判断。这些判断只能在动手写一个完整项目的过程中慢慢形成。'''
    )


def main():
    if not os.path.exists(TEMPLATE):
        raise SystemExit("找不到模板：" + TEMPLATE)
    doc = docx.Document(TEMPLATE)
    fill(doc)
    doc.save(OUT_DOCX)
    print("已生成：", OUT_DOCX)

    # ---- 校验：重新打开，确认结构与内容真的写进去了 ----
    check = docx.Document(OUT_DOCX)
    heads = [p.text for p in check.paragraphs if p.style.name == "Heading 1"]
    chars = sum(len(p.text) for p in check.paragraphs)
    print("标题：", heads)
    print("段落数=%d  表格数=%d  图片数=%d  正文字符数≈%d"
          % (len(check.paragraphs), len(check.tables), len(check.inline_shapes), chars))
    cover = [p.text.strip() for p in check.paragraphs[:18] if p.text.strip()]
    print("封面：", cover)

    shutil.copyfile(OUT_DOCX, DESKTOP_COPY)
    print("已复制到桌面：", DESKTOP_COPY)


if __name__ == "__main__":
    main()